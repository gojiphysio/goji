"""Build a structured XLSX template for clinicians to log sessions and
auto-compute commission. Replaces the free-text sheet that caused the
$0.90 typo we caught against the operations dashboard.

The Catalog is sourced from scripts/data/inventory.xlsx — a copy of the
Plato inventory export. When inventory changes (price tweaks, new SKUs,
deactivations), overwrite that file and rerun this generator.

Three sheets in the output workbook:
  - Log      Entry grid. Service picked from a dropdown sourced from Catalog.
             Per-session value auto-fills via VLOOKUP. Optional Override
             column handles patient-specific pricing (e.g. Medisave-discounted
             consults).
  - Catalog  Pulled from Plato inventory. Pack per-session values are
             FORMULAS so editing a pack price re-flows everywhere.
  - Summary  Pay-period rollup and commission calc. Base pay / threshold /
             rate are parametrised inputs.
"""

from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parent.parent
INVENTORY_PATH = REPO_ROOT / "scripts" / "data" / "inventory.xlsx"
OUT_PATH = Path.home() / "Downloads" / "Claire Revenue Tracker - Template.xlsx"

LOG_ROWS = 100

# Items in the inventory we DON'T track for clinician revenue:
# - Consumable (resistance bands, lyra) — sold not delivered
# - Misc (shaker bottles, vouchers) — merchandise / gift cards
TRACKED_TYPES = {"Consultation", "Package"}


def categorise(item_code: str, name: str) -> str:
    """Map a Plato item code to a clinical category for rollups."""
    c = item_code.replace(" ", "").upper()
    if c.startswith("PHY") or c == "CORPPT":
        return "Physiotherapy"
    if c.startswith("MB-"):
        return "Mummy & Baby"
    if c.startswith("CP"):
        return "Clinical Pilates"
    if c.startswith("PP"):
        return "Private Pilates"
    if c.startswith("SC-"):
        return "Strength & Conditioning"
    return "Other"


def category_sort_key(cat: str) -> int:
    order = {
        "Physiotherapy": 0,
        "Mummy & Baby": 1,
        "Clinical Pilates": 2,
        "Private Pilates": 3,
        "Strength & Conditioning": 4,
        "Other": 9,
    }
    return order.get(cat, 9)


def type_sort_key(typ: str) -> int:
    # Single-session consults before package redemptions inside each category
    return {"Consultation": 0, "Package": 1}.get(typ, 9)


def load_inventory() -> list[dict]:
    """Read scripts/data/inventory.xlsx and return rows we want to track."""
    wb = openpyxl.load_workbook(INVENTORY_PATH, data_only=True)
    ws = wb.active

    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    def get(row, key):
        i = idx.get(key)
        return row[i] if i is not None and i < len(row) else None

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        item_type = get(r, "Type")
        if item_type not in TRACKED_TYPES:
            continue
        if str(get(r, "Deactivate Item (All Locations)") or "FALSE").upper() == "TRUE":
            continue

        code = (get(r, "Inventory ID/Code") or "").strip()
        name = " ".join((get(r, "Name") or "").split())  # collapse double-spaces
        price = float(get(r, "Selling Price") or 0)
        redemptions = int(float(get(r, "Redemptions") or 0))

        if item_type == "Package":
            pack_size = redemptions if redemptions > 0 else 1
            display_name = f"REDEEM {name}"
        else:
            pack_size = 1
            display_name = name

        rows.append({
            "code": code,
            "name": name,
            "display": display_name,
            "type": item_type,
            "category": categorise(code, name),
            "pack_size": pack_size,
            "pack_price": price,
        })

    rows.sort(key=lambda x: (
        category_sort_key(x["category"]),
        type_sort_key(x["type"]),
        x["pack_size"],
        x["code"],
    ))
    return rows


HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LABEL_FONT = Font(bold=True, size=11)
SECTION_FILL = PatternFill("solid", fgColor="F3F4F6")
TOTAL_FILL = PatternFill("solid", fgColor="FEF3C7")
MUTED_FONT = Font(color="6B7280")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CURRENCY = '"$"#,##0.00'


def style_header(cell):
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = BORDER


def build_catalog(ws, catalog):
    headers = ["Item code", "Service / Redemption", "Category", "Type",
               "Pack size", "Pack price", "Per-session value", "Notes"]
    widths = [14, 64, 24, 14, 11, 14, 18, 50]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        style_header(c)
        ws.column_dimensions[get_column_letter(col)].width = w

    for i, item in enumerate(catalog, start=2):
        ws.cell(row=i, column=1, value=item["code"]).border = BORDER
        ws.cell(row=i, column=2, value=item["display"]).border = BORDER
        ws.cell(row=i, column=3, value=item["category"]).border = BORDER
        ws.cell(row=i, column=4, value=item["type"]).border = BORDER

        ps = ws.cell(row=i, column=5, value=item["pack_size"])
        ps.border = BORDER
        ps.alignment = Alignment(horizontal="right")

        pp = ws.cell(row=i, column=6, value=item["pack_price"])
        pp.number_format = CURRENCY
        pp.border = BORDER

        per = ws.cell(row=i, column=7, value=f"=F{i}/E{i}")
        per.number_format = CURRENCY
        per.font = Font(bold=True)
        per.border = BORDER

        ws.cell(row=i, column=8).border = BORDER

    ws.freeze_panes = "A2"


def build_log(ws, catalog_rows: int):
    headers = ["#", "Date", "Patient", "Service", "Per-session value",
               "Override value", "Final value", "Notes"]
    widths = [6, 14, 18, 64, 18, 18, 18, 36]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        style_header(c)
        ws.column_dimensions[get_column_letter(col)].width = w

    # Two hidden helper columns. Both kept on the SAME sheet for cross-tool
    # compat:
    #   I = Category, per-row VLOOKUP. Lets Summary roll up with plain
    #       SUMIFS instead of array-VLOOKUP-in-SUMPRODUCT (which Google
    #       Sheets evaluates as scalar, breaking the rollup).
    #   J = Dropdown source list, mirroring Catalog!B. Apple Numbers drops
    #       cross-sheet data-validation lists; same-sheet survives.
    cat_lookup_full = f"Catalog!$B$2:$G${catalog_rows + 1}"
    cat_lookup_bc = f"Catalog!$B$2:$C${catalog_rows + 1}"

    # Helper I: Category
    cat_header = ws.cell(row=1, column=9, value="(category)")
    cat_header.font = MUTED_FONT
    ws.column_dimensions["I"].width = 24
    ws.column_dimensions["I"].hidden = True

    # Helper J: Dropdown source
    src_header = ws.cell(row=1, column=10, value="(dropdown source)")
    src_header.font = MUTED_FONT
    ws.column_dimensions["J"].width = 50
    ws.column_dimensions["J"].hidden = True
    for i in range(catalog_rows):
        ws.cell(row=2 + i, column=10, value=f"=Catalog!B{2 + i}")

    same_sheet_list = f"$J$2:$J${catalog_rows + 1}"
    dv = DataValidation(
        type="list",
        formula1=f"={same_sheet_list}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Pick from catalog",
        error="Choose a service from the dropdown. To add a new service, edit the Catalog sheet first.",
    )
    dv.add(f"D2:D{LOG_ROWS + 1}")
    ws.add_data_validation(dv)

    # Date validation on column B — triggers the date picker UI in Google
    # Sheets after xlsx import. Range bounded to avoid Excel's "must be a
    # date" pedantry but wide enough to never matter in practice.
    date_dv = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(1900,1,1)",
        formula2="DATE(2200,12,31)",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Enter a valid date",
        error="Type a date or use the picker. Example: 2026-05-04",
    )
    date_dv.add(f"B2:B{LOG_ROWS + 1}")
    ws.add_data_validation(date_dv)

    for r in range(2, LOG_ROWS + 2):
        ws.cell(row=r, column=1, value=r - 1).border = BORDER
        ws.cell(row=r, column=2).number_format = "ddd, mmm d"
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=3).border = BORDER
        ws.cell(row=r, column=4).border = BORDER

        std = ws.cell(
            row=r, column=5,
            value=f'=IFERROR(VLOOKUP(D{r},{cat_lookup_full},6,FALSE),"")',
        )
        std.number_format = CURRENCY
        std.border = BORDER
        std.font = MUTED_FONT

        ovr = ws.cell(row=r, column=6)
        ovr.number_format = CURRENCY
        ovr.border = BORDER

        final = ws.cell(
            row=r, column=7,
            value=f'=IF(D{r}="","",IF(ISNUMBER(F{r}),F{r},IFERROR(E{r},0)))',
        )
        final.number_format = CURRENCY
        final.font = Font(bold=True)
        final.border = BORDER

        ws.cell(row=r, column=8).border = BORDER

        # Helper I: per-row category lookup (drives Summary breakdown)
        ws.cell(
            row=r, column=9,
            value=f'=IFERROR(VLOOKUP(D{r},{cat_lookup_bc},2,FALSE),"")',
        )

    ws.freeze_panes = "B2"


def build_summary(ws, log_rows: int, catalog):
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 16

    ws.cell(row=1, column=1, value="Commission summary").font = Font(bold=True, size=16)
    ws.cell(row=2, column=1, value="Pay-period dates default to the current month — edit them below.").font = Font(color="6B7280", italic=True)

    ws.cell(row=4, column=1, value="Settings").font = LABEL_FONT
    ws.cell(row=4, column=1).fill = SECTION_FILL

    # Default pay period = current month (live formulas, recompute on open).
    # Without these, B5/B6 are blank and SUMIFS criteria ">="&"" match no
    # dates, so every rollup shows 0.
    settings = [
        ("Pay period start", "=EOMONTH(TODAY(),-1)+1", "date"),
        ("Pay period end", "=EOMONTH(TODAY(),0)", "date"),
        ("Base pay", 5000.00, "currency"),
        ("Commission threshold", 13000.00, "currency"),
        ("Commission rate", 0.40, "percent"),
    ]
    for i, (label, default, fmt) in enumerate(settings, start=5):
        ws.cell(row=i, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=i, column=2, value=default)
        cell.border = BORDER
        if fmt == "currency":
            cell.number_format = CURRENCY
        elif fmt == "percent":
            cell.number_format = "0%"
        elif fmt == "date":
            cell.number_format = "ddd, mmm d, yyyy"

    # Date picker on pay-period cells (B5, B6)
    summary_date_dv = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(1900,1,1)",
        formula2="DATE(2200,12,31)",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Enter a valid date",
        error="Type a date or use the picker.",
    )
    summary_date_dv.add("B5:B6")
    ws.add_data_validation(summary_date_dv)

    base_pay_ref = "$B$7"
    threshold_ref = "$B$8"
    rate_ref = "$B$9"
    start_ref = "$B$5"
    end_ref = "$B$6"

    ws.cell(row=11, column=1, value="Calculation").font = LABEL_FONT
    ws.cell(row=11, column=1).fill = SECTION_FILL

    final_range = f"Log!$G$2:$G${log_rows + 1}"
    date_range = f"Log!$B$2:$B${log_rows + 1}"
    service_range = f"Log!$D$2:$D${log_rows + 1}"

    sessions_formula = (
        f'=COUNTIFS({final_range},">0",'
        f'{date_range},">="&{start_ref},'
        f'{date_range},"<="&{end_ref})'
    )
    revenue_formula = (
        f'=SUMIFS({final_range},'
        f'{date_range},">="&{start_ref},'
        f'{date_range},"<="&{end_ref})'
    )

    rows = [
        ("Sessions delivered", sessions_formula, "0"),
        ("Actual revenue (pay period)", revenue_formula, CURRENCY),
        ("Revenue above threshold", f"=MAX(0,B13-{threshold_ref})", CURRENCY),
        ("Commission earned", f"=B14*{rate_ref}", CURRENCY),
        ("Base pay", f"={base_pay_ref}", CURRENCY),
    ]
    for i, (label, formula, fmt) in enumerate(rows, start=12):
        ws.cell(row=i, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=i, column=2, value=formula)
        cell.number_format = fmt
        cell.border = BORDER

    ws.cell(row=17, column=1, value="Total payout").font = Font(bold=True, size=13)
    payout = ws.cell(row=17, column=2, value="=B15+B16")
    payout.number_format = CURRENCY
    payout.font = Font(bold=True, size=13)
    payout.fill = TOTAL_FILL
    payout.border = BORDER

    ws.cell(row=18, column=1, value="commission = max(0, revenue − threshold) × rate").font = Font(color="6B7280", italic=True, size=10)

    # Breakdown by category
    ws.cell(row=4, column=4, value="Breakdown by category").font = LABEL_FONT
    ws.cell(row=4, column=4).fill = SECTION_FILL
    for col, h in enumerate(["Category", "Sessions", "Revenue"], start=4):
        style_header(ws.cell(row=5, column=col, value=h))

    seen = []
    for item in catalog:
        if item["category"] not in seen:
            seen.append(item["category"])

    # Log!I is the per-row Category helper column. Plain SUMIFS / COUNTIFS
    # against it works identically in Excel, Google Sheets, and Numbers —
    # no array-VLOOKUP gymnastics required.
    category_range = f"Log!$I$2:$I${log_rows + 1}"

    for i, category in enumerate(seen, start=6):
        ws.cell(row=i, column=4, value=category).border = BORDER
        cnt_formula = (
            f'=COUNTIFS({category_range},D{i},'
            f'{date_range},">="&{start_ref},'
            f'{date_range},"<="&{end_ref})'
        )
        rev_formula = (
            f'=SUMIFS({final_range},'
            f'{category_range},D{i},'
            f'{date_range},">="&{start_ref},'
            f'{date_range},"<="&{end_ref})'
        )
        cnt = ws.cell(row=i, column=5, value=cnt_formula)
        cnt.number_format = "0"
        cnt.border = BORDER
        rev = ws.cell(row=i, column=6, value=rev_formula)
        rev.number_format = CURRENCY
        rev.border = BORDER

    total_row = 6 + len(seen)
    ws.cell(row=total_row, column=4, value="Total").font = LABEL_FONT
    ws.cell(row=total_row, column=4).fill = TOTAL_FILL
    ws.cell(row=total_row, column=4).border = BORDER
    cnt_total = ws.cell(row=total_row, column=5, value=f"=SUM(E6:E{total_row - 1})")
    cnt_total.number_format = "0"
    cnt_total.font = LABEL_FONT
    cnt_total.fill = TOTAL_FILL
    cnt_total.border = BORDER
    rev_total = ws.cell(row=total_row, column=6, value=f"=SUM(F6:F{total_row - 1})")
    rev_total.number_format = CURRENCY
    rev_total.font = LABEL_FONT
    rev_total.fill = TOTAL_FILL
    rev_total.border = BORDER

    ws.freeze_panes = "A4"


def patch_xlsx_for_google_sheets(path: Path):
    """openpyxl writes data validations with showDropDown="0" — in the OOXML
    spec 0 means "show the arrow", but Google Sheets reads the attribute
    INVERTED and hides the dropdown arrow, making the validation invisible
    to users. Excel reads it correctly either way. Fix: strip the attribute
    so Sheets falls back to the default (show arrow).

    Also strips the leading "=" some Sheets versions choke on inside
    <formula1>.

    See: https://issuetracker.google.com/issues/35949021 (Sheets DV import)
    """
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        unpacked = tmp / "unpacked"
        with zipfile.ZipFile(path, "r") as z:
            z.extractall(unpacked)

        for sheet_xml in (unpacked / "xl" / "worksheets").glob("sheet*.xml"):
            text = sheet_xml.read_text(encoding="utf-8")
            new_text = re.sub(r'\s*showDropDown="[01]"', '', text)
            new_text = re.sub(r'(<formula1>)=', r'\1', new_text)
            if new_text != text:
                sheet_xml.write_text(new_text, encoding="utf-8")

        repacked = tmp / "out.xlsx"
        with zipfile.ZipFile(repacked, "w", zipfile.ZIP_DEFLATED) as z:
            for f in unpacked.rglob("*"):
                if f.is_file():
                    z.write(f, f.relative_to(unpacked))
        shutil.copy(repacked, path)


def main():
    catalog = load_inventory()

    wb = Workbook()
    log = wb.active
    log.title = "Log"
    catalog_ws = wb.create_sheet("Catalog")
    summary_ws = wb.create_sheet("Summary")

    build_catalog(catalog_ws, catalog)
    build_log(log, catalog_rows=len(catalog))
    build_summary(summary_ws, log_rows=LOG_ROWS, catalog=catalog)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    patch_xlsx_for_google_sheets(OUT_PATH)
    print(f"Wrote {OUT_PATH}")
    print(f"  Catalog: {len(catalog)} services loaded from {INVENTORY_PATH.relative_to(REPO_ROOT)}")
    by_cat = {}
    for item in catalog:
        by_cat.setdefault(item["category"], []).append(item)
    for cat, items in sorted(by_cat.items(), key=lambda kv: category_sort_key(kv[0])):
        print(f"    {cat}: {len(items)}")


if __name__ == "__main__":
    main()
