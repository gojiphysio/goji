#!/usr/bin/env python3
"""
Row-level audit export of UNTYPED appointments (blank code_Bottom), for the
clinic to fix at source in Plato.

Two sheets:
  1. "Patient appts missing type" - real patient bookings with no type set.
     Patient shown as PII-minimised "F. Lastname" (from marts.dim_patient.
     initial_lastname - the raw name never reaches BigQuery). Auditable: the
     clinic can find each by date + time + clinician + initial.
  2. "Blocks - no patient" - untyped entries with no patient_id (admin blocks,
     leave, lunch, etc.). No name; these likely shouldn't count toward the
     treatment-mix denominator at all.

Output lands in reports/ which is gitignored - it carries patient initials +
dates, so it must never be committed. Forward it directly instead.

Usage:
    python3 scripts/build_untyped_audit_xlsx.py
"""
from __future__ import annotations

import json
import subprocess
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CLINICIAN_NAMES = {
    "PCL": "Cheryl Long (director)",
    "PKN": "Kristy Ng",
    "PCN": "Claire Ng",
}

# appt times are stored as naive local wall-clock parsed as UTC, so formatting
# the timestamp in UTC reproduces the clinic's original booking time.
PATIENT_SQL = """
SELECT
  DATE(a.start_at)                                  AS appt_date,
  FORMAT_TIMESTAMP('%H:%M', a.start_at, 'UTC')      AS appt_time,
  COALESCE(a.clinician_id, '(unmapped)')            AS clinician_id,
  COALESCE(a.calendar_color, '(none)')              AS calendar_color,
  p.initial_lastname                                AS patient
FROM `goji-bi.stg.appointments` a
JOIN `goji-bi.marts.dim_patient` p USING (patient_id)
WHERE a.appt_date >= '2026-01-01' AND a.treatment_type = 'Untyped'
ORDER BY appt_date, appt_time
"""

BLOCK_SQL = """
SELECT
  DATE(a.start_at)                                  AS appt_date,
  FORMAT_TIMESTAMP('%H:%M', a.start_at, 'UTC')      AS appt_time,
  COALESCE(a.clinician_id, '(unmapped)')            AS clinician_id,
  COALESCE(a.calendar_color, '(none)')              AS calendar_color
FROM `goji-bi.stg.appointments` a
WHERE a.appt_date >= '2026-01-01'
  AND a.treatment_type = 'Untyped'
  AND (a.patient_id IS NULL OR a.patient_id = '')
ORDER BY appt_date, appt_time
"""

NAVY = "1F3864"
LIGHT = "D9E1F2"
GREY = "808080"


def bq_json(sql: str) -> list[dict]:
    # bq query defaults to 100 rows; raise the cap so large result sets aren't
    # silently truncated.
    out = subprocess.run(
        ["bq", "query", "--use_legacy_sql=false", "--max_rows=100000",
         "--format=json", sql],
        check=True, capture_output=True, text=True,
    ).stdout
    return json.loads(out)


def clinician(code: str) -> str:
    return CLINICIAN_NAMES.get(code, f"Unknown calendar ({code})")


def style_header(ws, headers: list[str], row: int) -> None:
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def write_rows(ws, rows: list[list], start_row: int, name_cols: set[int]) -> int:
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    r = start_row
    for values in rows:
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if c in name_cols else "center",
                vertical="center")
        r += 1
    return r


def build(patient_rows: list[dict], block_rows: list[dict], dest: Path) -> None:
    wb = Workbook()

    # --- Sheet 1: patient appts missing a type (the actionable list) ---
    ws = wb.active
    ws.title = "Patient appts missing type"
    ws["A1"] = "Goji Health - patient appointments with NO treatment type set"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.merge_cells("A1:E1")
    ws["A2"] = (
        f"2026 year-to-date. Generated {date.today().isoformat()}. "
        f"{len(patient_rows)} appointments. "
        "Each is a real patient booking where the type (S&C / Pilates / Physio / "
        "Home Visit) was left blank in Plato. Please set the type on each."
    )
    ws["A3"] = (
        "CONFIDENTIAL - contains patient initials. Patient shown as first "
        "initial + last name. Find each in Plato by date + time + clinician."
    )
    for rr in (2, 3):
        ws[f"A{rr}"].font = Font(italic=True, color=GREY, size=10)
        ws[f"A{rr}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(f"A{rr}:E{rr}")
    ws.row_dimensions[2].height = 30

    headers = ["Date", "Time", "Clinician", "Calendar", "Patient (initial + last name)"]
    style_header(ws, headers, 5)
    data = [[r["appt_date"], r["appt_time"], clinician(r["clinician_id"]),
             r["calendar_color"], r["patient"]] for r in patient_rows]
    write_rows(ws, data, 6, name_cols={3, 5})
    for i, w in enumerate([12, 8, 26, 12, 30], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"

    # --- Sheet 2: blocks / admin time (no patient) ---
    ws2 = wb.create_sheet("Blocks - no patient")
    ws2["A1"] = "Untyped entries with NO patient attached (blocks / admin / leave)"
    ws2["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws2.merge_cells("A1:D1")
    ws2["A2"] = (
        f"{len(block_rows)} entries. These have no patient_id - they are calendar "
        "blocks, not treatments, so they arguably should not count in the "
        "treatment-mix chart. Listed for completeness; no action needed unless "
        "any are actually mis-booked patient sessions."
    )
    ws2["A2"].font = Font(italic=True, color=GREY, size=10)
    ws2["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws2.merge_cells("A2:D2")
    ws2.row_dimensions[2].height = 30
    style_header(ws2, ["Date", "Time", "Clinician", "Calendar"], 4)
    bdata = [[r["appt_date"], r["appt_time"], clinician(r["clinician_id"]),
              r["calendar_color"]] for r in block_rows]
    write_rows(ws2, bdata, 5, name_cols={3})
    for i, w in enumerate([12, 8, 26, 12], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A5"

    wb.save(dest)


def main() -> None:
    patient_rows = bq_json(PATIENT_SQL)
    block_rows = bq_json(BLOCK_SQL)
    dest = (Path(__file__).resolve().parents[1] / "reports"
            / "Goji - untyped appointments audit (2026 YTD).xlsx")
    dest.parent.mkdir(exist_ok=True)
    build(patient_rows, block_rows, dest)
    print(f"wrote {dest}")
    print(f"  patient appts missing type: {len(patient_rows)}")
    print(f"  blocks (no patient):        {len(block_rows)}")


if __name__ == "__main__":
    main()
